import docx

def read_docx(filename):
    doc = docx.Document(filename)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    
    print("\n--- TABLES ---")
    for table in doc.tables:
        for row in table.rows:
            print(" | ".join([cell.text.replace("\n", " ") for cell in row.cells]))
        print("-" * 20)
        
    return '\n'.join(fullText)

print(read_docx("Claim Component I August 2025.docx"))
